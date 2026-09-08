from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import markdown
from bs4 import BeautifulSoup
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import validate_distribution_tables as validator
from export_sample_docs import md_to_html_body


RAW = "{保留第一行说明}\n{保留第二行说明}\n\n{生命=勇气×10+等级}\n"


class RawBraceCoverageTests(unittest.TestCase):
    def test_raw_oracle_detects_attr_list_loss_even_when_rendered_prose_omits_it(self):
        blocks = validator.raw_brace_blocks(RAW, "车卡.md")
        broken = BeautifulSoup(markdown.markdown(RAW, extensions=["attr_list"]), "html.parser")
        coverage = validator.raw_brace_coverage(blocks, {"车卡.md": validator.visible_html_text(broken)})
        self.assertEqual(len(blocks), 3)
        self.assertEqual([block["line"] for block in coverage["missing_blocks"]], [2])

    def test_current_export_parser_preserves_consecutive_braces_and_formula(self):
        with tempfile.TemporaryDirectory() as raw:
            source = Path(raw) / "车卡.md"
            source.write_text(RAW, encoding="utf-8")
            soup = BeautifulSoup(md_to_html_body(source), "html.parser")
        coverage = validator.raw_brace_coverage(validator.raw_brace_blocks(RAW, "车卡.md"),
                                                {"车卡.md": validator.visible_html_text(soup)})
        self.assertEqual(coverage["matched_block_count"], 3)
        self.assertFalse(coverage["missing_blocks"])

    def test_another_source_section_cannot_satisfy_missing_brace(self):
        blocks = validator.raw_brace_blocks("{相同的核心说明}", "车卡.md")
        html = '<section><p class="source">车卡.md</p><p>丢失</p></section>'
        html += '<section><p class="source">技能.md</p><p>{相同的核心说明}</p></section>'
        coverage = validator.raw_brace_coverage(blocks, validator.html_section_texts(BeautifulSoup(html, "html.parser")))
        self.assertEqual(len(coverage["missing_blocks"]), 1)

    def test_missing_repeated_occurrence_within_source_is_detected(self):
        blocks = validator.raw_brace_blocks("{重复提示}\n{重复提示}", "车卡.md")
        coverage = validator.raw_brace_coverage(blocks, {"车卡.md": "{重复提示}"})
        self.assertEqual(coverage["matched_block_count"], 1)
        self.assertEqual(coverage["missing_blocks"][0]["required_occurrence"], 2)

    def test_html_attributes_styles_and_scripts_are_not_visible_proof(self):
        blocks = validator.raw_brace_blocks("{中文提示}", "车卡.md")
        html = '<section><p class="source">车卡.md</p><p title="中文提示">正文</p>'
        html += '<style>/* 中文提示 */</style><script>"中文提示"</script></section>'
        coverage = validator.raw_brace_coverage(blocks, validator.html_section_texts(BeautifulSoup(html, "html.parser")))
        self.assertEqual(coverage["matched_block_count"], 0)

    def test_pdf_markers_bound_local_content_despite_line_wrapping(self):
        files = ["核心/技能.md", "森语/车卡.md", "森语/战斗.md"]
        text = "核心/技能.md {核心同句} 森语/车\n卡.md {生命 = 勇气×10} 森语/战斗.md {核心同句}"
        sections = validator.text_sections_by_source_markers(text, files)
        blocks = validator.raw_brace_blocks("{核心同句}\n{生命=勇气×10}", files[1])
        coverage = validator.raw_brace_coverage(blocks, sections)
        self.assertEqual(coverage["matched_block_count"], 1)
        self.assertEqual(coverage["missing_blocks"][0]["text"], "{核心同句}")

    def test_docx_visible_text_is_checked_in_own_source_section(self):
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "sample.docx"
            document = Document()
            document.add_paragraph("甲.md", style="Caption")
            document.add_paragraph("{同一说明}")
            document.add_paragraph("乙.md", style="Caption")
            document.add_paragraph("丢失")
            document.save(path)
            summary = validator.docx_summary(path, validator.raw_brace_blocks("{同一说明}", "乙.md"), ["甲.md", "乙.md"])
        self.assertEqual(summary["raw_brace_coverage"]["matched_block_count"], 0)

    def test_removing_attr_list_keeps_explicit_html_attributes_and_table_geometry(self):
        markup = '<table data-table-id="fixture" class="pmtrpg-table"><tr><td rowspan="2" colspan="2">甲</td><td>乙</td></tr><tr><td>丙</td></tr></table>\n\n<a id="anchor" href="#anchor">链接</a>'
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            path = root / "sample.md"
            path.write_text(markup, encoding="utf-8")
            soup = BeautifulSoup(md_to_html_body(path), "html.parser")
            with patch.object(validator, "DOCS_DIR", root):
                summary = validator.html_summary([path])
        self.assertEqual(soup.table["data-table-id"], "fixture")
        self.assertIn("pmtrpg-table", soup.table["class"])
        self.assertEqual(soup.a["href"], "#anchor")
        self.assertEqual(soup.a["id"], "anchor")
        self.assertEqual(summary["table_structures"][0]["merged_cells"], [
            {"row": 0, "col": 0, "rowspan": 2, "colspan": 2}])
        self.assertEqual(summary["docx_gridspan_slots"], 2)


if __name__ == "__main__":
    unittest.main()
