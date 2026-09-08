from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from export_sample_docs import (  # noqa: E402
    CHM_CONTROL_ENCODING,
    _add_docx_table,
    _cell_text_without_nested_tables,
    _direct_nested_tables,
    _inline_html,
    _local_path_from_uri,
    _table_layout,
    build_chm_name_map,
    build_hhc,
)
from audit_export_readiness import analyze_file, bucket_file  # noqa: E402


class ExportTableAdapterTests(unittest.TestCase):
    def parse_table(self, source: str):
        return BeautifulSoup(source, "html.parser").find("table")

    def test_layout_reads_sectioned_rows_and_spans(self) -> None:
        table = self.parse_table(
            """
            <table>
              <thead><tr><th colspan="3">标题</th></tr></thead>
              <tbody>
                <tr><td rowspan="2">A</td><td>B</td><td>C</td></tr>
                <tr><td colspan="2">D</td></tr>
              </tbody>
            </table>
            """
        )
        anchors, rows, columns, header_rows = _table_layout(table)
        self.assertEqual((rows, columns, header_rows), (3, 3, 1))
        self.assertEqual(
            [(item["row"], item["col"], item["rowspan"], item["colspan"]) for item in anchors],
            [(0, 0, 1, 3), (1, 0, 2, 1), (1, 1, 1, 1), (1, 2, 1, 1), (2, 1, 1, 2)],
        )

    def test_nested_table_is_not_counted_as_outer_rows(self) -> None:
        table = self.parse_table(
            """
            <table><tr><td>外层<table><tr><td>内层</td></tr></table></td></tr></table>
            """
        )
        _anchors, rows, columns, _header_rows = _table_layout(table)
        self.assertEqual((rows, columns), (1, 1))
        cell = table.find("td")
        self.assertEqual(len(_direct_nested_tables(cell)), 1)

    def test_docx_adapter_emits_horizontal_and_vertical_merges(self) -> None:
        table_node = self.parse_table(
            """
            <table>
              <tr><td colspan="2">横向</td></tr>
              <tr><td rowspan="2">纵向</td><td>一</td></tr>
              <tr><td>二</td></tr>
            </table>
            """
        )
        document = Document()
        _add_docx_table(document, table_node)
        xml = document._element.xml
        self.assertIn("w:gridSpan", xml)
        self.assertIn("w:vMerge", xml)

    def test_combined_merge_repeats_gridspan_for_vertical_continuations(self) -> None:
        table_node = self.parse_table(
            '<table><tr><td colspan="2" rowspan="3">A</td><td>B</td></tr>'
            '<tr><td>C</td></tr><tr><td>D</td></tr></table>'
        )
        document = Document()
        _add_docx_table(document, table_node)
        self.assertEqual(len(document._element.xpath('.//w:gridSpan')), 3)
        self.assertEqual(len(document._element.xpath('.//w:vMerge')), 3)

    def test_windows_file_uri_round_trip_handles_percent_encoding(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "带 空格.png"
            path.write_bytes(b"image")
            self.assertEqual(_local_path_from_uri(path.resolve().as_uri()), path.resolve())

    def test_existing_cell_images_do_not_emit_duplicate_alt_text(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            image_path = Path(raw) / "黑桃.png"
            Image.new("RGBA", (8, 8), (0, 0, 0, 255)).save(image_path)
            table = self.parse_table(
                f'<table><tr><td>视觉<img src="{image_path.resolve().as_uri()}" alt="黑桃"></td></tr></table>'
            )
            cell = table.find("td")
            self.assertEqual(_cell_text_without_nested_tables(cell), "视觉")

            document = Document()
            result = _add_docx_table(document, table)
            exported_cell = result.cell(0, 0)
            self.assertEqual(len(exported_cell.paragraphs), 1)
            self.assertNotIn("[黑桃]", exported_cell.text)
            self.assertIn("w:drawing", exported_cell._tc.xml)

    def test_inline_skill_marker_keeps_source_order_in_docx_and_pdf_markup(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            image_path = Path(raw) / "梅花.png"
            Image.new("RGBA", (8, 8), (0, 0, 0, 255)).save(image_path)
            table = self.parse_table(
                f'<table><tr><td>手术<img src="{image_path.resolve().as_uri()}" alt="梅花">：说明</td></tr></table>'
            )
            cell = table.find("td")
            pdf_markup = _inline_html(cell, remove_nested_tables=True)
            self.assertLess(pdf_markup.index("手术"), pdf_markup.index("<img"))
            self.assertLess(pdf_markup.index("<img"), pdf_markup.index("：说明"))
            self.assertNotIn("〔双倍〕", pdf_markup)

            document = Document()
            result = _add_docx_table(document, table)
            xml = result.cell(0, 0)._tc.xml
            self.assertLess(xml.index("手术"), xml.index("w:drawing"))
            self.assertLess(xml.index("w:drawing"), xml.index("：说明"))

    def test_chm_toc_round_trips_through_legacy_control_encoding(self) -> None:
        entries = [("核心规则/速查图表/技能列表.md", "技能列表")]
        payload = build_hhc(entries, build_chm_name_map(entries))
        encoded = payload.encode(CHM_CONTROL_ENCODING)
        self.assertIn("技能列表", encoded.decode(CHM_CONTROL_ENCODING))
        self.assertIn("charset=gb2312", payload)

    def test_export_readiness_counts_nested_canonical_tables_without_blocking(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "sample.md"
            path.write_text(
                "# 示例\n<table class=\"pmtrpg-table\" data-table-id=\"outer\"><tr><td>"
                "<table class=\"pmtrpg-table\" data-table-id=\"inner\"><tr><td>x</td></tr></table>"
                "</td></tr></table>\n",
                encoding="utf-8",
            )
            info = analyze_file(path, "sample.md")
            self.assertEqual(info["tables"], 2)
            self.assertEqual(info["canonical_tables"], 2)
            self.assertEqual(info["noncanonical_tables"], 0)
            self.assertNotEqual(bucket_file(info), "A-阻塞导出")

    def test_export_readiness_blocks_noncanonical_table(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            path = Path(raw) / "sample.md"
            path.write_text("# 示例\n<table><tr><td>x</td></tr></table>\n", encoding="utf-8")
            info = analyze_file(path, "sample.md")
            self.assertEqual(info["noncanonical_tables"], 1)
            self.assertEqual(bucket_file(info), "A-阻塞导出")


if __name__ == "__main__":
    unittest.main()
