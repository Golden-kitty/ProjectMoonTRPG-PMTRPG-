from __future__ import annotations

import argparse
import hashlib
import html
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from urllib.parse import unquote, urlparse, urlunparse

import markdown
from bs4 import BeautifulSoup, NavigableString

REPO_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = REPO_ROOT / "docs"
OUTPUT_DIR = REPO_ROOT / "output" / "export_samples"
CHM_CONTROL_ENCODING = "gbk"

DEFAULT_SAMPLES: list[tuple[str, str]] = [
    ("核心规则/基本规则/等级.md", "等级"),
    ("核心规则/速查图表/技能列表.md", "技能列表"),
    ("核心规则/战斗/战斗流程.md", "战斗流程"),
    ("资源目录/装备/武器/奇门.md", "奇门"),
]


def md_to_html_body(md_path: Path) -> str:
    # Canonical tables are already explicit in the source. Export must not
    # infer or rewrite merges as a hidden rendering step.
    text = md_path.read_text(encoding="utf-8")
    md = markdown.Markdown(
        # Rule prose uses literal {Chinese notes}. attr_list would silently
        # consume a note on the final line of a paragraph as HTML attributes.
        # Canonical tables already carry their explicit HTML attributes.
        extensions=["tables", "md_in_html", "sane_lists", "toc"]
    )
    body = md.convert(text)

    def repl(match: re.Match[str]) -> str:
        attr = match.group(1)
        raw = match.group(2)
        if re.match(r"^[a-z]+://", raw, re.I) or raw.startswith("data:") or raw.startswith("#"):
            return match.group(0)
        parsed = urlparse(html.unescape(raw))
        resolved = (md_path.parent / unquote(parsed.path)).resolve()
        local = urlparse(resolved.as_uri())
        target = urlunparse(local._replace(query=parsed.query, fragment=parsed.fragment))
        return f'{attr}="{html.escape(target, quote=True)}"'

    body = re.sub(r'(src|href)="([^"]+)"', repl, body)
    body = re.sub(r"\n{3,}", "\n\n", body)
    return body


def _is_table_intro(node) -> bool:
    """Keep short labels with the following block, not preceding prose."""
    text = node.get_text(strip=True)
    following = node.find_next_sibling()
    return (
        node.find_parent("table") is None
        and 0 < len(text) <= 30
        and re.search(r"[。！？；，,.!?;]", text) is None
        and following is not None
        and (
            following.name in {"table", "h1", "h2", "h3", "h4", "h5", "h6", "hr"}
            or (following.name == "p" and re.fullmatch(r"【[^】]{1,12}】", text) is not None)
        )
    )


def build_full_html(title: str, sections: list[tuple[str, str]]) -> str:
    section_map = {
        (DOCS_DIR / rel_path).resolve().as_uri(): index
        for index, (rel_path, _label) in enumerate(sections, start=1)
    }
    toc = "\n".join(
        f'<li><a href="#section-{idx}">{html.escape(label)}</a></li>'
        for idx, (_path, label) in enumerate(sections, start=1)
    )
    section_html = []
    for idx, (rel_path, label) in enumerate(sections, start=1):
        body = md_to_html_body(DOCS_DIR / rel_path)
        body_soup = BeautifulSoup(body, "html.parser")
        for node in body_soup.find_all(id=True):
            node["id"] = f'section-{idx}--{node["id"]}'
        for link in body_soup.find_all("a", href=True):
            parsed = urlparse(link["href"])
            target_index = idx if link["href"].startswith("#") else section_map.get(
                urlunparse(parsed._replace(query="", fragment=""))
            )
            if target_index is not None:
                suffix = f"--{parsed.fragment}" if parsed.fragment else ""
                link["href"] = f"#section-{target_index}{suffix}"
        for icon in body_soup.find_all("img", src=True):
            if "技能列表.files" in unquote(icon["src"]):
                icon["class"] = [*icon.get("class", []), "pmtrpg-icon"]
        for table in body_soup.find_all("table"):
            # Wide numeric tables need room for all columns and the outer
            # border inside the printable content box.
            if _table_layout(table)[2] >= 16:
                table["class"] = [*table.get("class", []), "pmtrpg-wide-table"]
            # Keep short cards/templates together, but permit long or nested
            # tables to split naturally instead of clipping them to a page.
            if not table.find("table") and len(table.find_all("tr")) <= 10 and len(table.get_text(strip=True)) <= 600:
                table["class"] = [*table.get("class", []), "pmtrpg-short-table"]
        for paragraph in body_soup.find_all("p"):
            if _is_table_intro(paragraph):
                paragraph["class"] = [*paragraph.get("class", []), "pmtrpg-table-intro"]
        body = str(body_soup)
        section_html.append(
            "\n".join(
                [
                    f'<section id="section-{idx}">',
                    f"<h1>{html.escape(label)}</h1>",
                    f'<p class="source">{html.escape(rel_path)}</p>',
                    body,
                    "</section>",
                ]
            )
        )
    return "\n".join(
        [
            "<!doctype html>",
            '<html lang="zh-CN">',
            "<head>",
            '  <meta charset="utf-8">',
            f"  <title>{html.escape(title)}</title>",
            "  <style>",
            "    @page { size: A4; margin: 12mm; }",
            "    body { font-family: 'Microsoft YaHei', sans-serif; line-height: 1.6; margin: 24px; }",
            "    h1, h2, h3, h4, h5, h6 { page-break-after: avoid; break-after: avoid; }",
            "    section { page-break-before: always; }",
            "    section:first-of-type { page-break-before: auto; }",
            "    table { border-collapse: collapse; width: 100%; font-size: 12px; }",
            "    th, td { border: 1px solid #999; padding: 4px 6px; vertical-align: top; }",
            "    table.pmtrpg-wide-table { font-size: 11px; }",
            "    table.pmtrpg-wide-table > tbody > tr > td, table.pmtrpg-wide-table > tbody > tr > th, table.pmtrpg-wide-table > tr > td, table.pmtrpg-wide-table > tr > th, table.pmtrpg-wide-table > thead > tr > th, table.pmtrpg-wide-table > thead > tr > td { padding: 3px 2px; }",
            "    .pmtrpg-table-intro { break-after: avoid; page-break-after: avoid; }",
            "    hr { break-after: avoid; page-break-after: avoid; }",
            "    thead { display: table-header-group; }",
            "    img { max-width: 100%; height: auto; }",
            "    img.pmtrpg-icon { width: 1em; height: 1em; vertical-align: text-bottom; }",
            "    img[src*='技能列表.files'] { width: 1em; height: 1em; vertical-align: text-bottom; }",
            "    .source { color: #666; font-size: 12px; }",
            "    @media print { body { margin: 0; } table.pmtrpg-short-table { break-inside: avoid; page-break-inside: avoid; } }",
            "  </style>",
            "</head>",
            "<body>",
            f"  <h1>{html.escape(title)}</h1>",
            "  <p>本文件为 PMTRPG 规则与表格的分发审阅稿。</p>",
            "  <h2>目录</h2>",
            f"  <ul>{toc}</ul>",
            *section_html,
            "</body>",
            "</html>",
            "",
        ]
    )


def ensure_clean_dir(path: Path) -> None:
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True, exist_ok=True)


def parse_stub_policy(policy_path: Path | None) -> set[str]:
    if not policy_path or not policy_path.exists():
        return set()

    filtered: set[str] = set()
    current = ""
    for raw in policy_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if line.startswith("## "):
            current = line[3:].strip().lower()
            continue
        if not line.startswith("- `"):
            continue
        if current not in {"导出过滤", "export filter"}:
            continue
        filtered.add(line.split("`", 2)[1].replace("docs/", ""))
    return filtered


def build_sections(
    files: list[str] | None,
    file_list: Path | None,
    stub_policy: Path | None,
) -> list[tuple[str, str]]:
    rel_paths: list[str]
    if files:
        rel_paths = [path.replace("\\", "/").removeprefix("docs/") for path in files]
    elif file_list:
        rel_paths = []
        for raw in file_list.read_text(encoding="utf-8").splitlines():
            line = raw.strip()
            if not line or line.startswith("#"):
                continue
            rel_paths.append(line.replace("\\", "/").removeprefix("docs/"))
    else:
        rel_paths = [path for path, _label in DEFAULT_SAMPLES]

    filtered = parse_stub_policy(stub_policy)
    rel_paths = [path for path in rel_paths if path not in filtered]
    return [(path, Path(path).stem) for path in rel_paths]


def batch_paths(batch_name: str) -> tuple[Path, str]:
    if batch_name == "sample-export":
        return OUTPUT_DIR, "PMTRPG Export Samples"
    return OUTPUT_DIR / batch_name, f"PMTRPG Export Batch: {batch_name}"


def write_verification_record(
    out_dir: Path,
    batch_name: str,
    sections: list[tuple[str, str]],
    generated: list[Path],
) -> Path:
    record_path = out_dir / "verification.md"
    lines = [
        f"# {batch_name}",
        "",
        "## Files",
        "",
        *[f"- `{rel_path}`" for rel_path, _label in sections],
        "",
        "## Generated",
        "",
        *[f"- `{path.relative_to(REPO_ROOT).as_posix()}`" for path in generated if path.exists()],
        "",
    ]
    record_path.write_text("\n".join(lines), encoding="utf-8")
    return record_path


def build_html_preview(out_dir: Path, sections: list[tuple[str, str]], title: str, batch_name: str) -> Path:
    html_path = out_dir / f"{batch_name}.html"
    html_path.write_text(build_full_html(title, sections), encoding="utf-8")
    return html_path


def _inline_html(node, *, remove_nested_tables: bool = False) -> str:
    node = BeautifulSoup(str(node), "html.parser")
    root = node.find()
    if root is None:
        return ""
    if remove_nested_tables:
        for nested in root.find_all("table"):
            nested.decompose()
    for img in node.find_all("img"):
        src = img.get("src", "")
        local_path = _local_path_from_uri(src)
        if local_path is not None:
            img.attrs = {
                "src": local_path.resolve().as_posix(),
                "width": "9",
                "height": "9",
                "valign": "middle",
            }
        else:
            alt = img.get("alt") or Path(src.replace("file:///", "")).name or "image"
            img.replace_with(node.new_string(f"[图:{alt}]"))
    text = root.decode_contents()
    text = re.sub(r"</?(?:p|div|section|tbody|thead|tfoot)>", "", text)
    text = " ".join(text.split())
    return text.strip()


def _local_path_from_uri(value: str) -> Path | None:
    """Resolve local file URLs emitted by md_to_html_body on Windows."""
    if not value.startswith("file:"):
        return None
    parsed = urlparse(value)
    raw_path = unquote(parsed.path)
    if re.match(r"^/[A-Za-z]:/", raw_path):
        raw_path = raw_path[1:]
    path = Path(raw_path)
    return path if path.exists() else None


def _direct_table_rows(table_node) -> list:
    return [
        row
        for row in table_node.find_all("tr")
        if row.find_parent("table") is table_node
    ]


def _direct_row_cells(row_node) -> list:
    return [
        cell
        for cell in row_node.find_all(["th", "td"], recursive=False)
        if cell.find_parent("tr") is row_node
    ]


def _positive_span(cell, name: str) -> int:
    try:
        return max(1, int(cell.get(name, 1)))
    except (TypeError, ValueError):
        return 1


def _table_layout(table_node) -> tuple[list[dict], int, int, int]:
    """Return anchor cells plus the logical grid dimensions for an HTML table."""
    rows = _direct_table_rows(table_node)
    occupied: set[tuple[int, int]] = set()
    anchors: list[dict] = []
    max_col = 0
    header_rows = 0

    for row_index, row in enumerate(rows):
        col_index = 0
        if row.find_parent("thead") is not None:
            header_rows = row_index + 1
        for cell in _direct_row_cells(row):
            while (row_index, col_index) in occupied:
                col_index += 1
            row_span = _positive_span(cell, "rowspan")
            col_span = _positive_span(cell, "colspan")
            anchors.append(
                {
                    "row": row_index,
                    "col": col_index,
                    "rowspan": row_span,
                    "colspan": col_span,
                    "node": cell,
                    "is_header": cell.name.lower() == "th",
                }
            )
            for rr in range(row_index, row_index + row_span):
                for cc in range(col_index, col_index + col_span):
                    occupied.add((rr, cc))
            col_index += col_span
            max_col = max(max_col, col_index)

    max_row = max(
        [len(rows), *[item["row"] + item["rowspan"] for item in anchors]],
        default=0,
    )
    return anchors, max_row, max_col, header_rows


def _cell_text_without_nested_tables(cell_node) -> str:
    clone = BeautifulSoup(str(cell_node), "html.parser")
    clone_cell = clone.find(["th", "td"])
    if clone_cell is None:
        return ""
    for nested in clone_cell.find_all("table"):
        nested.decompose()
    for image in clone_cell.find_all("img"):
        src = image.get("src", "")
        if _local_path_from_uri(src) is not None:
            image.decompose()
        else:
            alt = image.get("alt") or Path(src or "image").stem
            image.replace_with(clone.new_string(f"[{alt}]"))
    for br in clone_cell.find_all("br"):
        br.replace_with(clone.new_string(" "))
    return " ".join(str(clone_cell.decode_contents()).split())


def _direct_cell_images(cell_node) -> list[Path]:
    owner_table = cell_node.find_parent("table")
    paths: list[Path] = []
    for image in cell_node.find_all("img"):
        if image.find_parent("table") is not owner_table:
            continue
        path = _local_path_from_uri(image.get("src", ""))
        if path is not None:
            paths.append(path)
    return paths


def _direct_nested_tables(cell_node) -> list:
    owner_table = cell_node.find_parent("table")
    return [
        table
        for table in cell_node.find_all("table")
        if table.find_parent("table") is owner_table
    ]


def _append_docx_inline(paragraph, node, *, skip_nested_tables: bool = False) -> None:
    """Append mixed text, line breaks, and local images in source order."""

    def visit(current) -> None:
        if isinstance(current, NavigableString):
            value = re.sub(r"\s+", " ", str(current))
            if value.strip():
                paragraph.add_run(value)
            return
        tag = getattr(current, "name", "").lower()
        if not tag:
            return
        if tag == "table" and skip_nested_tables:
            return
        if tag == "br":
            paragraph.add_run().add_break()
            return
        if tag == "img":
            src = current.get("src", "")
            image_path = _local_path_from_uri(src)
            if image_path is not None:
                _add_docx_image(paragraph._parent, image_path, paragraph=paragraph)
            else:
                alt = current.get("alt") or Path(src or "image").stem
                paragraph.add_run(f"[图:{alt}]")
            return

        start = len(paragraph.runs)
        for child in current.children:
            visit(child)
        if tag in {"b", "strong", "i", "em"}:
            for run in paragraph.runs[start:]:
                if tag in {"b", "strong"}:
                    run.bold = True
                if tag in {"i", "em"}:
                    run.italic = True

    for child in node.children:
        visit(child)


def _add_images_from_node(node, flowables, max_width: int = 420) -> None:
    images = [node] if getattr(node, "name", "").lower() == "img" else node.find_all("img")
    for img in images:
        src = img.get("src", "")
        local_path = _local_path_from_uri(src)
        if local_path is None:
            continue
        try:
            from reportlab.lib import utils
            from reportlab.platypus import Image, Spacer

            reader = utils.ImageReader(str(local_path))
            width, height = reader.getSize()
            scale = min(1.0, max_width / width) if width else 1.0
            flowables.append(Image(str(local_path), width=width * scale, height=height * scale))
            flowables.append(Spacer(1, 6))
        except Exception:
            continue


def _reportlab_table_from_html(table_node, base, colors, max_width: float = 515):
    from reportlab.lib import utils
    from reportlab.platypus import Image, Paragraph, Table, TableStyle

    anchors, row_count, col_count, header_rows = _table_layout(table_node)
    if not anchors or row_count == 0 or col_count == 0:
        return None

    data: list[list[object]] = [
        [Paragraph(" ", base) for _ in range(col_count)] for _ in range(row_count)
    ]
    styles: list[tuple] = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]

    for anchor in anchors:
        row = anchor["row"]
        col = anchor["col"]
        node = anchor["node"]
        content: list[object] = []
        text = _inline_html(node, remove_nested_tables=True)
        if text:
            content.append(Paragraph(text, base))
        nested_width = max_width * anchor["colspan"] / col_count - 8
        for nested in _direct_nested_tables(node):
            nested_table = _reportlab_table_from_html(nested, base, colors, max(60, nested_width))
            if nested_table is not None:
                content.append(nested_table)
        data[row][col] = content or Paragraph(" ", base)

        end_col = col + anchor["colspan"] - 1
        end_row = row + anchor["rowspan"] - 1
        if end_col != col or end_row != row:
            styles.append(("SPAN", (col, row), (end_col, end_row)))
        if anchor["is_header"]:
            styles.append(("BACKGROUND", (col, row), (end_col, end_row), colors.whitesmoke))

    col_widths = [max_width / col_count] * col_count
    table = Table(data, colWidths=col_widths, repeatRows=header_rows, splitByRow=1)
    table.setStyle(TableStyle(styles))
    return table


def _html_sections_to_pdf_flowables(sections: list[tuple[str, str]]):
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.pdfbase.pdfmetrics import registerFont
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:  # pragma: no cover
        raise SystemExit(
            "reportlab is required for PDF export. Install it with "
            "`python -m pip install -r requirements-export.txt`."
        ) from exc

    configured_font = os.environ.get("PMTRPG_PDF_FONT")
    font_candidates = [
        Path(configured_font).expanduser() if configured_font else None,
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/System/Library/Fonts/PingFang.ttc"),
    ]
    font_path = next((path for path in font_candidates if path is not None and path.is_file()), None)
    if font_path is None:
        raise SystemExit(
            "PDF export requires an embeddable CJK font. Set PMTRPG_PDF_FONT to a TTF/TTC file."
        )
    registerFont(TTFont("PmChinese", str(font_path), subfontIndex=0))

    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "PmBase",
        parent=styles["BodyText"],
        fontName="PmChinese",
        fontSize=10,
        leading=14,
        alignment=TA_LEFT,
    )
    h1 = ParagraphStyle("PmH1", parent=base, fontSize=18, leading=22, spaceAfter=10)
    h2 = ParagraphStyle("PmH2", parent=base, fontSize=14, leading=18, spaceAfter=8)
    meta = ParagraphStyle("PmMeta", parent=base, fontSize=8, leading=10, textColor=colors.grey)

    flowables = [Paragraph("PMTRPG Export Samples", h1), Paragraph("本文件用于 PMTRPG 导出就绪样本验证。", base), Spacer(1, 8)]

    def append_node(node) -> None:
        if isinstance(node, NavigableString):
            text = str(node).strip()
            if text:
                flowables.append(Paragraph(html.escape(text), base))
                flowables.append(Spacer(1, 4))
            return
        if not getattr(node, "name", None):
            return

        tag = node.name.lower()
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            text = _inline_html(node)
            if text:
                flowables.append(Paragraph(text, h2))
            return
        if tag in {"p", "blockquote", "pre", "dt", "dd", "summary"}:
            text = _inline_html(node)
            if text:
                flowables.append(Paragraph(text, base))
                flowables.append(Spacer(1, 4))
            return
        if tag in {"ul", "ol"}:
            for li in node.find_all("li", recursive=False):
                text = _inline_html(li)
                if text:
                    flowables.append(Paragraph(f"• {text}", base))
            flowables.append(Spacer(1, 4))
            return
        if tag == "table":
            table = _reportlab_table_from_html(node, base, colors)
            if table is not None:
                flowables.append(table)
                flowables.append(Spacer(1, 8))
            return
        if tag == "img":
            _add_images_from_node(node, flowables)
            return
        if tag == "hr":
            flowables.append(Spacer(1, 8))
            return

        # Markdown-in-HTML and raw source fragments commonly introduce div,
        # section, details, or project-specific wrappers. Recurse rather than
        # silently dropping their visible text.
        for child in node.children:
            append_node(child)

    for idx, (rel_path, label) in enumerate(sections, start=1):
        if idx > 1:
            flowables.append(PageBreak())
        flowables.append(Paragraph(label, h1))
        flowables.append(Paragraph(rel_path, meta))
        flowables.append(Spacer(1, 8))

        body = md_to_html_body(DOCS_DIR / rel_path)
        soup = BeautifulSoup(f"<root>{body}</root>", "html.parser")
        for node in soup.root.children:
            append_node(node)

    return flowables


def _find_chromium() -> Path | None:
    configured = os.environ.get("PMTRPG_CHROMIUM")
    candidates = [
        Path(configured).expanduser() if configured else None,
        *(Path(found) for name in ("chrome", "msedge", "chromium") if (found := shutil.which(name))),
        Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
        Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
    ]
    return next((path for path in candidates if path is not None and path.is_file()), None)


def export_pdf(out_dir: Path, sections: list[tuple[str, str]], title: str, batch_name: str) -> Path:
    html_path = build_html_preview(out_dir, sections, title, batch_name)
    pdf_path = out_dir / f"{batch_name}.pdf"
    chromium = _find_chromium()
    if chromium is None:
        raise SystemExit(
            "PDF export requires Chromium. Set PMTRPG_CHROMIUM to chrome.exe, msedge.exe, or chromium."
        )
    with tempfile.TemporaryDirectory(prefix="pmtrpg-pdf-") as profile_dir:
        result = subprocess.run(
            [
                str(chromium),
                "--headless=new",
                "--disable-gpu",
                "--allow-file-access-from-files",
                "--no-pdf-header-footer",
                f"--user-data-dir={profile_dir}",
                f"--print-to-pdf={pdf_path.resolve()}",
                html_path.resolve().as_uri(),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=90,
        )
    if result.returncode != 0 or not pdf_path.is_file() or pdf_path.stat().st_size == 0:
        detail = (result.stderr or result.stdout).strip()
        raise SystemExit(f"Chromium PDF export failed ({result.returncode}): {detail}")
    return pdf_path


def _set_repeat_table_header(row) -> None:
    from docx.oxml import OxmlElement

    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "true")
    tr_pr.append(marker)


def _compact_docx_paragraph(paragraph) -> None:
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.line_spacing = 1


def _add_docx_image(container, image_path: Path, max_inches: float = 1.2, paragraph=None) -> None:
    from docx.shared import Inches
    from PIL import Image as PilImage

    try:
        with PilImage.open(image_path) as source:
            width, height = source.size
        target = min(max_inches, 0.17 if max(width, height) <= 128 else max_inches)
        target_paragraph = paragraph if paragraph is not None else container.add_paragraph()
        _compact_docx_paragraph(target_paragraph)
        run = target_paragraph.add_run()
        picture = run.add_picture(str(image_path), width=Inches(target))
        for doc_pr in picture._inline.xpath(".//wp:docPr"):
            doc_pr.set("descr", image_path.name)
            doc_pr.set("title", image_path.stem)
    except Exception:
        target_paragraph = paragraph if paragraph is not None else container.add_paragraph()
        target_paragraph.add_run(f"[图:{image_path.name}]")


def _add_docx_table(container, table_node):
    anchors, row_count, col_count, header_rows = _table_layout(table_node)
    if not anchors or row_count == 0 or col_count == 0:
        return None

    table = container.add_table(rows=row_count, cols=col_count)
    table.style = "Table Grid"
    for anchor in anchors:
        row = anchor["row"]
        col = anchor["col"]
        end_row = row + anchor["rowspan"] - 1
        end_col = col + anchor["colspan"] - 1
        cell = table.cell(row, col)
        if end_row != row or end_col != col:
            cell = cell.merge(table.cell(end_row, end_col))
        cell.text = ""
        cell_paragraph = cell.paragraphs[0]
        _compact_docx_paragraph(cell_paragraph)
        _append_docx_inline(cell_paragraph, anchor["node"], skip_nested_tables=True)
        for nested in _direct_nested_tables(anchor["node"]):
            _add_docx_table(cell, nested)

    for row in table.rows[:header_rows]:
        _set_repeat_table_header(row)
    return table


def _compact_docx_table_paragraphs(table) -> None:
    seen_cells: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            marker = id(cell._tc)
            if marker in seen_cells:
                continue
            seen_cells.add(marker)
            for paragraph in cell.paragraphs:
                _compact_docx_paragraph(paragraph)
            for nested in cell.tables:
                _compact_docx_table_paragraphs(nested)


def export_docx(out_dir: Path, sections: list[tuple[str, str]], title: str, batch_name: str) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt

    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / f"{batch_name}.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1
    document.styles["Caption"].paragraph_format.keep_with_next = True
    for level in range(1, 7):
        document.styles[f"Heading {level}"].paragraph_format.keep_with_next = True
        document.styles[f"Heading {level}"].paragraph_format.keep_together = True

    def append_node(node) -> None:
        if isinstance(node, NavigableString):
            text = str(node).strip()
            if text:
                paragraph = document.add_paragraph(text)
                _compact_docx_paragraph(paragraph)
            return
        if not getattr(node, "name", None):
            return

        tag = node.name.lower()
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            level = min(6, int(tag[1]))
            document.add_heading(node.get_text(" ", strip=True), level=level)
            return
        if tag in {"p", "blockquote", "pre", "dt", "dd", "summary"}:
            paragraph = document.add_paragraph()
            _append_docx_inline(paragraph, node)
            if _is_table_intro(node):
                paragraph.paragraph_format.keep_with_next = True
            return
        if tag in {"ul", "ol"}:
            style = "List Bullet" if tag == "ul" else "List Number"
            for item in node.find_all("li", recursive=False):
                paragraph = document.add_paragraph(style=style)
                _append_docx_inline(paragraph, item)
            return
        if tag == "table":
            _add_docx_table(document, node)
            spacer = document.add_paragraph()
            _compact_docx_paragraph(spacer)
            return
        if tag == "img":
            image_path = _local_path_from_uri(node.get("src", ""))
            if image_path is not None:
                _add_docx_image(document, image_path, max_inches=5.5)
            return
        if tag == "hr":
            return
        for child in node.children:
            append_node(child)

    document.add_heading(title, level=0)
    document.add_paragraph("本文件用于 PMTRPG canonical HTML 表格分发适配验证。")
    for _index, (rel_path, label) in enumerate(sections):
        document.add_heading(label, level=1)
        document.add_paragraph(rel_path, style="Caption")
        body = md_to_html_body(DOCS_DIR / rel_path)
        soup = BeautifulSoup(f"<root>{body}</root>", "html.parser")
        for node in soup.root.children:
            append_node(node)

    for table in document.tables:
        _compact_docx_table_paragraphs(table)

    document.save(docx_path)
    return docx_path


def build_chm_name_map(entries: list[tuple[str, str]]) -> dict[str, str]:
    return {
        rel_path: f"sample-{index:02d}.html"
        for index, (rel_path, _label) in enumerate(entries, start=1)
    }


def build_hhc(entries: list[tuple[str, str]], name_map: dict[str, str]) -> str:
    items = []
    for rel_path, label in entries:
        html_name = name_map[rel_path]
        items.extend(
            [
                "<LI> <OBJECT type=\"text/sitemap\">",
                f"    <param name=\"Name\" value=\"{html.escape(label)}\">",
                f"    <param name=\"Local\" value=\"html/{html_name}\">",
                "  </OBJECT>",
            ]
        )
    return "\n".join(
        [
            "<!DOCTYPE HTML PUBLIC \"-//IETF//DTD HTML//EN\">",
            "<HTML>",
            "<HEAD><META http-equiv=\"Content-Type\" content=\"text/html; charset=gb2312\"></HEAD>",
            "<BODY>",
            "<UL>",
            *items,
            "</UL>",
            "</BODY>",
            "</HTML>",
            "",
        ]
    )


def export_chm(out_dir: Path, sections: list[tuple[str, str]], title: str, batch_name: str, compile_chm: bool) -> Path:
    chm_dir = out_dir / "chm"
    html_dir = chm_dir / "html"
    ensure_clean_dir(chm_dir)
    html_dir.mkdir(parents=True, exist_ok=True)

    name_map = build_chm_name_map(sections)
    source_name_map = {(DOCS_DIR / path).resolve().as_uri(): name for path, name in name_map.items()}
    html_files: list[str] = []
    asset_dir = chm_dir / "assets"
    asset_dir.mkdir(parents=True, exist_ok=True)
    asset_files: dict[Path, str] = {}
    for rel_path, label in sections:
        html_name = name_map[rel_path]
        html_files.append(f"html/{html_name}")
        page_html = build_full_html(f"{label} - PMTRPG Export Sample", [(rel_path, label)])
        page_soup = BeautifulSoup(page_html, "html.parser")
        for link in page_soup.find_all("a", href=True):
            parsed = urlparse(link["href"])
            source_key = urlunparse(parsed._replace(query="", fragment=""))
            if source_key in source_name_map:
                suffix = f"#section-1--{parsed.fragment}" if parsed.fragment else ""
                link["href"] = source_name_map[source_key] + suffix
        for image in page_soup.find_all("img"):
            source = _local_path_from_uri(image.get("src", ""))
            if source is None:
                continue
            if source not in asset_files:
                digest = hashlib.sha256(source.read_bytes()).hexdigest()[:12]
                target_name = f"asset-{digest}{source.suffix.lower()}"
                shutil.copy2(source, asset_dir / target_name)
                asset_files[source] = target_name
            image["src"] = f"../assets/{asset_files[source]}"
        page_html = str(page_soup)
        (html_dir / html_name).write_text(page_html, encoding="utf-8")

    hhc_path = chm_dir / f"{batch_name}.hhc"
    # HTML Help Workshop 4.x reads project/control files through the active
    # Windows ANSI code page. UTF-8 makes Chinese TOC labels mojibake even when
    # the UTF-8 topic HTML itself renders correctly.
    hhc_path.write_text(build_hhc(sections, name_map), encoding=CHM_CONTROL_ENCODING)

    hhp_path = chm_dir / f"{batch_name}.hhp"
    hhp_path.write_text(
        "\n".join(
            [
                "[OPTIONS]",
                "Compatibility=1.1 or later",
                f"Compiled file={batch_name}.chm",
                f"Contents file={batch_name}.hhc",
                f"Default topic=html/{name_map[sections[0][0]]}",
                "Display compile progress=No",
                "Full-text search=Yes",
                "Language=0x804 Chinese (Simplified, PRC)",
                f"Title={title}",
                "",
                "[FILES]",
                *html_files,
                *[f"assets/{name}" for name in sorted(set(asset_files.values()))],
                "",
            ]
        ),
        encoding=CHM_CONTROL_ENCODING,
    )

    hhc_exe = shutil.which("hhc")
    if compile_chm and hhc_exe:
        subprocess.run([hhc_exe, str(hhp_path)], cwd=chm_dir, check=False)
    return hhp_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", choices=["pdf", "chm", "docx", "all"], default="all")
    parser.add_argument("--compile-chm", action="store_true")
    parser.add_argument("--batch-name", default="sample-export")
    parser.add_argument("--title", help="Human-readable book or review title")
    parser.add_argument("--files", nargs="*")
    parser.add_argument("--file-list", type=Path)
    parser.add_argument("--stub-policy", type=Path, default=REPO_ROOT / "docs" / "acceptance" / "export-stub-page-policy.md")
    args = parser.parse_args()

    sections = build_sections(args.files, args.file_list, args.stub_policy)
    if not sections:
        raise SystemExit("No exportable files selected.")

    out_dir, title = batch_paths(args.batch_name)
    title = args.title or title
    out_dir.mkdir(parents=True, exist_ok=True)
    generated: list[Path] = []

    if args.format in {"pdf", "all"}:
        pdf_dir = out_dir / "pdf"
        ensure_clean_dir(pdf_dir)
        pdf_path = export_pdf(pdf_dir, sections, title, args.batch_name)
        generated.extend([pdf_path, pdf_dir / f"{args.batch_name}.html"])
        print(f"pdf {pdf_path}")

    if args.format in {"docx", "all"}:
        docx_dir = out_dir / "docx"
        ensure_clean_dir(docx_dir)
        docx_path = export_docx(docx_dir, sections, title, args.batch_name)
        generated.append(docx_path)
        print(f"docx {docx_path}")

    if args.format in {"chm", "all"}:
        chm_path = export_chm(out_dir, sections, title, args.batch_name, compile_chm=args.compile_chm)
        generated.extend(
            [
                chm_path,
                chm_path.with_suffix(".hhc"),
                chm_path.with_suffix(".chm"),
            ]
        )
        print(f"chm {chm_path}")
        if args.compile_chm and not shutil.which("hhc"):
            print("chm_compile_skipped missing_hhc")

    record_path = write_verification_record(out_dir, args.batch_name, sections, generated)
    print(f"record {record_path}")


if __name__ == "__main__":
    main()
